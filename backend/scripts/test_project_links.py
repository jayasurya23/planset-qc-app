"""Project IDs: a QC project's link to PMO 360 and monday.com.

The Castillo Project ID ("264-066", "2512-053" -- named after the "Project ID"
column on monday.com's Portfolio board) is the key PMO 360, monday and the
drawings' title block share. These tests pin:

  normalisation   pasted values (en dashes, stray spaces) compare equal, and
                  the value is otherwise left opaque.
  monday client   read-only (a mutation is refused before sending), errors
                  inside an HTTP 200 are errors, rate limits retry, columns are
                  found by TITLE, pages are followed, the Portfolio is cached.
  resolution      linked / not_found / ambiguous / error / not_configured, with
                  the Project ID saved in every case, and a lookup that
                  finishes after it changed cannot attach the old project's
                  board.
  title block     the value is read from beside "CASTILLO PROJECT ID" on a
                  rotated production-geometry sheet, not from any number-shaped
                  string on the drawing.
  API             PATCH/GET projects, suggestions, the Portfolio picker, the
                  upload form field, and the Excel export.

No network: monday's transport is replaced with a fake board.

Run: PYTHONPATH=backend python backend/scripts/test_project_links.py
"""
import copy
import io
import os
import sys
import tempfile
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

os.environ["PLANSET_DATA_DIR"] = tempfile.mkdtemp(prefix="project_id_links_")
os.environ["DEV_USER_EMAIL"] = "engineer@castillope.com"
os.environ.pop("MONDAY_API_TOKEN", None)
os.environ.pop("PMO360_BASE_URL", None)

import fitz  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app import db, jobs, monday, project_links  # noqa: E402
from app.analyzer import detect_castillo_project_id  # noqa: E402
from app.main import app  # noqa: E402

_FAILS: list[str] = []


def check(name: str, cond: bool) -> None:
    print(f"  {'PASS' if cond else 'FAIL'}  {name}")
    if not cond:
        _FAILS.append(name)


# ── a fake monday.com ────────────────────────────────────────────────────
# Column ids are deliberately NOT the production ids: the client must find
# columns by title.
COLUMNS = [
    {"id": "name", "title": "Name"},
    {"id": "txt_pid_fake", "title": "Project ID"},
    {"id": "txt_client_fake", "title": "Client Name"},
    {"id": "status_fake", "title": "Contract Status"},
    {"id": "rel_fake", "title": "Project Tasks Links"},
]


def _item(iid, name, pid, client="Recon Corporation", status="Active"):
    return {"id": iid, "name": name, "column_values": [
        {"id": "txt_pid_fake", "text": pid},
        {"id": "txt_client_fake", "text": client},
        {"id": "status_fake", "text": status},
    ]}


FAKE = {
    "page1": [_item("101", "Aurora 1 IFC", "2512-053"),
              _item("102", "Highland North (1 & 2)", " 254 – 325 ", client="Montante Solar")],
    "page2": [_item("103", "Sawyer Road", ""), _item("104", "Coal City 2 IFC", "2512-061")],
    "links": {"101": ["9001", "9002"], "102": [], "104": ["9101"]},
    "boards": {"9001": ("777", "Aurora 1 IFC"), "9101": ("778", "Coal City 2 IFC")},
    "calls": [],
    "fail": None,           # None | "graphql" | "429-once" | "401" | "down"
}


def fake_post(url, payload, headers, timeout):
    q = payload["query"]
    FAKE["calls"].append(q)
    if FAKE["fail"] == "down":
        raise monday.MondayError("Could not reach monday.com (ConnectError).")
    if FAKE["fail"] == "401":
        return 401, {}, None
    if FAKE["fail"] == "429-once":
        FAKE["fail"] = None
        return 429, {"Retry-After": "0"}, None
    if FAKE["fail"] == "graphql":
        return 200, {}, {"errors": [{"message": "Column not found",
                                     "extensions": {"code": "InvalidColumnIdException"}}]}
    v = payload["variables"]
    if "next_items_page" in q:
        return 200, {}, {"data": {"next_items_page": {"cursor": None, "items": FAKE["page2"]}}}
    if "items_page" in q:
        return 200, {}, {"data": {"boards": [{"items_page": {"cursor": "c2", "items": FAKE["page1"]}}]}}
    if "BoardRelationValue" in q:
        iid = v["ids"][0]
        return 200, {}, {"data": {"items": [{"column_values": [
            {"linked_item_ids": FAKE["links"].get(iid, [])}]}]}}
    if "board { id name }" in q:
        tid = v["ids"][0]
        bid, bname = FAKE["boards"][tid]
        return 200, {}, {"data": {"items": [{"id": tid, "board": {"id": bid, "name": bname}}]}}
    if "columns" in q:
        return 200, {}, {"data": {"boards": [{"name": "Portfolio", "columns": COLUMNS}]}}
    raise AssertionError(f"unexpected query: {q[:80]}")


monday._post = fake_post
monday.time.sleep = lambda s: None


def configure(on: bool) -> None:
    if on:
        os.environ["MONDAY_API_TOKEN"] = "test-token"
    else:
        os.environ.pop("MONDAY_API_TOKEN", None)
    monday.reset_cache()
    FAKE["calls"].clear()
    FAKE["fail"] = None


# ── normalisation ────────────────────────────────────────────────────────
print("Project IDs are cleaned, not interpreted:")
n = monday.normalize_project_id
check("en dash and padding -> 264-066", n("  264–066 ") == "264-066")
check("spaces around the dash collapse", n("2512 - 053") == "2512-053")
check("em dash and minus sign are dashes too", n("2512—053") == n("2512−053") == "2512-053")
check("horizontal bar and fullwidth hyphen too -- the same set the TypeScript keys use",
      n("264―066") == n("264－066") == "264-066")
check("blank -> None", n("   ") is None and n(None) is None)
check("anything else is left alone", n("KA340 phase B") == "KA340 phase B")
check("comparison ignores case and dash style",
      monday.project_id_key("ka340-b") == monday.project_id_key("KA340–B"))

# ── the monday client ────────────────────────────────────────────────────
print("The monday client is read-only and honest about failure:")
configure(False)
try:
    monday.execute("query { me { id } }")
    check("no token -> MondayAuthError", False)
except monday.MondayAuthError:
    check("no token -> MondayAuthError", True)

configure(True)
try:
    monday.execute("mutation { delete_item(item_id: 1) { id } }")
    check("a mutation is refused", False)
except monday.MondayError:
    check("a mutation is refused", True)
check("...before anything is sent", FAKE["calls"] == [])

FAKE["fail"] = "graphql"
try:
    monday.execute("query { boards { id } }")
    check("GraphQL errors inside HTTP 200 raise", False)
except monday.MondayError as exc:
    check("GraphQL errors inside HTTP 200 raise", "Column not found" in str(exc))

configure(True)
FAKE["fail"] = "401"
try:
    monday.execute("query { boards { id } }")
    check("HTTP 401 -> MondayAuthError", False)
except monday.MondayAuthError:
    check("HTTP 401 -> MondayAuthError", True)

configure(True)
FAKE["fail"] = "429-once"
data = monday.execute("query ($ids: [ID!]) { boards(ids: $ids) { name columns { id title } } }",
                      {"ids": ["1"]})
check("a 429 is retried and then succeeds", bool(data.get("boards")) and len(FAKE["calls"]) == 2)

print("The Portfolio is read by column title, across pages, and cached:")
configure(True)
items = monday.portfolio()
by_name = {i["name"]: i for i in items}
check("both pages were read", len(items) == 4 and len(FAKE["calls"]) == 3)
check("columns resolved by title, not id",
      by_name["Aurora 1 IFC"]["castillo_project_id"] == "2512-053")
check("monday's own values are normalised",
      by_name["Highland North (1 & 2)"]["castillo_project_id"] == "254-325")
check("a blank Project ID is None", by_name["Sawyer Road"]["castillo_project_id"] is None)
check("client and status come along", by_name["Aurora 1 IFC"]["client"] == "Recon Corporation"
      and by_name["Aurora 1 IFC"]["status"] == "Active")
calls = len(FAKE["calls"])
monday.portfolio()
check("a second read within the TTL makes no call", len(FAKE["calls"]) == calls)
monday.portfolio(force=True)
check("force refetches the items", len(FAKE["calls"]) > calls)

print("A failing monday degrades without lying or stalling:")
FAKE["fail"] = "down"
try:
    monday.portfolio(force=True)
    check("a refresh that was asked for raises instead of answering from the old copy", False)
except monday.MondayError:
    check("a refresh that was asked for raises instead of answering from the old copy", True)
calls = len(FAKE["calls"])
check("during the back-off an automatic lookup gets the old copy without calling monday",
      len(monday.portfolio()) == 4 and len(FAKE["calls"]) == calls)
try:
    monday.portfolio(force=True)
    check("during the back-off a forced lookup fails at once", False)
except monday.MondayError:
    check("during the back-off a forced lookup fails at once", len(FAKE["calls"]) == calls)
FAKE["fail"] = None
check("an explicit refresh may go back to monday during the back-off",
      len(monday.portfolio(force=True, during_backoff=True)) == 4 and len(FAKE["calls"]) > calls)
os.environ["MONDAY_CACHE_TTL_SECONDS"] = "0"
FAKE["fail"] = "down"
check("an expired cache that cannot refresh still serves list callers",
      len(monday.portfolio()) == 4)
FAKE["fail"] = None
os.environ.pop("MONDAY_CACHE_TTL_SECONDS")
monday.reset_cache()

print("A recreated monday column is found again without a restart:")
configure(True)
monday.portfolio()
for it in FAKE["page1"] + FAKE["page2"]:
    for cv in it["column_values"]:
        if cv["id"] == "txt_pid_fake":
            cv["id"] = "txt_pid_new"
COLUMNS[1]["id"] = "txt_pid_new"
# monday answers an unknown column id with nothing, not an error.
_real_fake_post = fake_post


def dropping_post(url, payload, headers, timeout):
    status, hdrs, body = _real_fake_post(url, payload, headers, timeout)
    body = copy.deepcopy(body)          # never edit the fake board itself
    cols = (payload.get("variables") or {}).get("cols")
    if cols and isinstance(body, dict) and body.get("data"):
        pages = [b.get("items_page") for b in (body["data"].get("boards") or [])] + [
            body["data"].get("next_items_page")]
        for page in filter(None, pages):
            for it in page.get("items") or []:
                it["column_values"] = [cv for cv in it["column_values"] if cv["id"] in cols]
    return status, hdrs, body


monday._post = dropping_post
got = monday.portfolio(force=True)
check("a forced refresh re-reads column ids", sorted(i["castillo_project_id"] or "" for i in got)
      == ["", "2512-053", "2512-061", "254-325"])
monday.reset_cache()
monday._columns["18403099969"] = {"project id": "txt_pid_fake", "client name": "txt_client_fake",
                                  "contract status": "status_fake",
                                  "project tasks links": "rel_fake"}
got = monday.portfolio()
check("an ordinary fetch that reads every value blank re-reads column ids once",
      any(i["castillo_project_id"] == "2512-053" for i in got))
for it in FAKE["page1"] + FAKE["page2"]:
    for cv in it["column_values"]:
        if cv["id"] == "txt_pid_new":
            cv["id"] = "txt_pid_fake"
COLUMNS[1]["id"] = "txt_pid_fake"
monday._post = fake_post
configure(True)
check("find_by_project_id matches a pasted variant", [i["item_id"] for i in
      monday.find_by_project_id("254–325")] == ["102"])
check("the board comes from a linked task", monday.board_for_item("101") ==
      {"board_id": "777", "board_name": "Aurora 1 IFC"})
check("an item with no linked tasks has no board", monday.board_for_item("102") is None)


# ── title block ──────────────────────────────────────────────────────────
def rotated_sheet(entries):
    """Production geometry: authored portrait 1728x2592, shown landscape via
    /Rotate 270. Entries are DISPLAY coordinates; display (dx, dy) is authored
    (1728 - dy, dx)."""
    doc = fitz.open()
    page = doc.new_page(width=1728, height=2592)
    for dx, dy, txt, size in entries:
        page.insert_text((1728 - dy, dx), txt, fontsize=size, rotate=270)
    page.set_rotation(270)
    return doc


# Measured on real Castillo sheets: label at display (2337..2485, 631..646),
# value directly beneath at (2380..2472, 663..688).
TITLE_BLOCK = [(2337, 644, "CASTILLO PROJECT ID", 11), (2380, 686, "2512-053", 18)]
DECOYS = [
    (300, 900, "250-600", 14),                     # conductor size on the drawing
    (1200, 644, "UTILITY PROJECT ID", 11),         # somebody else's number, same layout
    (1200, 686, "175-225", 18),
]

print("The title block's Project ID is read beside its label:")
doc = rotated_sheet(DECOYS + TITLE_BLOCK)
check("rotated sheet with decoys -> 2512-053", detect_castillo_project_id(doc) == "2512-053")
doc.close()
doc = rotated_sheet(DECOYS)
check("number-shaped text without the label -> None", detect_castillo_project_id(doc) is None)
doc.close()
doc = fitz.open()
page = doc.new_page(width=2592, height=1728)
# The base-14 fonts cannot encode an en dash (it comes back as U+FFFD), so
# embed one that can -- real title blocks use real fonts.
page.insert_font(fontname="F0", fontbuffer=fitz.Font("cjk").buffer)
page.insert_text((2000, 1500), "CASTILLO PROJECT NO.", fontsize=11, fontname="F0")
page.insert_text((2160, 1500), "241–001", fontsize=11, fontname="F0")
check("older layout, value to the right, en dash -> 241-001",
      detect_castillo_project_id(doc) == "241-001")
doc.close()


# ── API ──────────────────────────────────────────────────────────────────
db.init_db()
RUNS_DIR = Path(os.environ["PLANSET_DATA_DIR"]) / "runs"


def seed_run(project_name, *, pdf_entries=None, detected=None):
    pid = db.get_or_create_project(project_name, "seed@castillope.com")
    rid = str(uuid.uuid4())
    run_dir = RUNS_DIR / rid
    run_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = run_dir / f"{project_name}.pdf"
    doc = rotated_sheet(pdf_entries or [])
    doc.save(pdf_path)
    doc.close()
    db.insert_run({
        "id": rid, "project_name": project_name, "original_filename": pdf_path.name,
        "created_at": "2026-09-01T00:00:00+00:00", "pdf_path": str(pdf_path),
        "page_count": 1, "project_id": pid,
        "summary": {"pdf_page_count": 1, "indexed_sheet_count": 1, "actual_sheet_count": 1,
                    "title_block_project_id": detected},
        "status_counts": {"Pass": 1}, "categories": [], "issues": [],
    }, [])
    return pid, rid


AURORA, AURORA_RUN = seed_run("Aurora 1 IFC 60% planset", pdf_entries=DECOYS + TITLE_BLOCK)
HIGHLAND, _ = seed_run("Highland North 1", detected="254-325")
jobs_submitted = []
jobs.submit = lambda job_id, kind, meta, fn: jobs_submitted.append(job_id)

# No token at startup: the startup refresh thread is tested directly below,
# not raced against these requests.
configure(False)
with TestClient(app) as client:
    print("Setting a Project ID links the project:")
    configure(True)
    r = client.patch(f"/api/projects/{AURORA}", json={"castillo_project_id": " 2512–053 "})
    p = r.json()
    check("PATCH -> 200", r.status_code == 200)
    check("stored normalised", p["castillo_project_id"] == "2512-053")
    check("attributed to the signed-in engineer", p["castillo_project_id_set_by"] == "engineer@castillope.com")
    check("monday status is linked", p["monday"]["status"] == "linked")
    check("monday board link", p["links"]["monday_board"] == "https://castillope.monday.com/boards/777")
    check("Portfolio item link", p["links"]["monday_item"] ==
          "https://castillope.monday.com/boards/18403099969/pulses/101")
    check("PMO 360 link by Project ID", p["links"]["pmo360"] ==
          "https://pmo360.castillope.com/portfolio?project_id=2512-053")
    check("no duplicates yet", p["duplicates"] == [])

    calls = len(FAKE["calls"])
    listed = {x["id"]: x for x in client.get("/api/projects").json()}
    check("listing carries Project ID and links", listed[AURORA]["links"]["monday_board"]
          == "https://castillope.monday.com/boards/777")
    check("listing never calls monday", len(FAKE["calls"]) == calls)

    print("Every monday outcome still saves the Project ID:")
    r = client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "254-325"})
    p = r.json()
    check("item without a board: linked, no board link, item link kept",
          p["monday"]["status"] == "linked" and p["links"]["monday_board"] is None
          and p["links"]["monday_item"] and "no project board" in (p["monday"]["detail"] or ""))

    r = client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "999-999"})
    p = r.json()
    check("unknown Project ID -> not_found, saved", p["monday"]["status"] == "not_found"
          and p["castillo_project_id"] == "999-999" and p["links"]["monday_board"] is None)
    check("changing the Project ID cleared the old board", db.get_project(HIGHLAND)["monday_board_id"] is None)

    FAKE["page2"].append(_item("105", "Coal City 2 IFC (copy)", "2512-061"))
    monday.reset_cache()
    r = client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "2512-061"})
    p = r.json()
    check("two Portfolio items share it -> ambiguous, nothing linked",
          p["monday"]["status"] == "ambiguous" and p["links"]["monday_board"] is None)
    FAKE["page2"].pop()
    monday.reset_cache()

    FAKE["fail"] = "graphql"
    r = client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "254-325"})
    p = r.json()
    check("monday failing -> 200, status error, Project ID saved",
          r.status_code == 200 and p["monday"]["status"] == "error" and p["castillo_project_id"] == "254-325")
    FAKE["fail"] = None
    r = client.post(f"/api/projects/{HIGHLAND}/monday-refresh")
    check("refresh recovers once monday answers", r.json()["monday"]["status"] == "linked")

    print("A shared Project ID warns rather than refuses:")
    r = client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "2512-053"})
    dup = r.json()["duplicates"]
    check("duplicate reported by id and name",
          dup == [{"id": AURORA, "name": "Aurora 1 IFC 60% planset"}])

    print("A lookup that outlives its Project ID cannot attach the wrong board:")
    stale = db.set_project_monday_link(
        HIGHLAND, castillo_project_id="999-000", status="linked", board_id="1")
    check("write for an old Project ID is ignored", stale is None
          and db.get_project(HIGHLAND)["monday_board_id"] != "1")

    print("Without a monday token the rest still works:")
    configure(False)
    r = client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "264-066"})
    p = r.json()
    check("not_configured, PMO 360 link still built",
          p["monday"]["status"] == "not_configured" and p["monday"]["configured"] is False
          and p["links"]["pmo360"].endswith("?project_id=264-066"))
    check("Portfolio picker reports not configured",
          client.get("/api/monday/portfolio").json() == {"configured": False, "items": [], "error": None})

    os.environ["PMO360_BASE_URL"] = ""
    check("PMO360_BASE_URL='' hides the PMO 360 link",
          client.get(f"/api/projects/{HIGHLAND}").json()["links"]["pmo360"] is None)
    os.environ["PMO360_BASE_URL"] = "https://pmo360-staging.example.io/"
    check("PMO360_BASE_URL points at another host",
          client.get(f"/api/projects/{HIGHLAND}").json()["links"]["pmo360"]
          == "https://pmo360-staging.example.io/portfolio?project_id=264-066")
    os.environ.pop("PMO360_BASE_URL")

    print("Turning monday on later re-checks the waiting projects in one pass:")
    check("bulk refresh without a token -> 409",
          client.post("/api/monday/refresh-projects").status_code == 409)
    configure(True)
    r = client.post("/api/monday/refresh-projects")
    check("only projects not yet linked are checked",
          r.status_code == 200 and r.json() == {"checked": 1, "by_status": {"not_found": 1}})
    check("the Portfolio is fetched once, not once per unmatched project",
          len(FAKE["calls"]) == 3)
    check("the waiting project now has a real status",
          client.get(f"/api/projects/{HIGHLAND}").json()["monday"]["status"] == "not_found")
    r = client.post("/api/monday/refresh-projects?include_linked=true")
    check("include_linked re-checks linked projects too",
          r.json() == {"checked": 2, "by_status": {"linked": 1, "not_found": 1}})
    configure(False)

    print("Clearing and rejecting:")
    r = client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "  "})
    p = r.json()
    check("blank clears Project ID, attribution and links", p["castillo_project_id"] is None
          and p["castillo_project_id_set_by"] is None and p["links"]["pmo360"] is None
          and p["monday"]["status"] is None)
    check("over 50 characters -> 422",
          client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "9" * 51}).status_code == 422)
    check("control characters -> 422",
          client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "12\x00-3"}).status_code == 422)
    check("unknown project -> 404",
          client.patch("/api/projects/nope", json={"castillo_project_id": "1-1"}).status_code == 404)
    check("a leading formula character -> 422 (it would run in the Excel export)",
          all(client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": v}).status_code == 422
              for v in ("=1+1", "+1", "-1", "@SUM(A1)")))

    print("Saved without a token, a Project ID is 'unchecked' once one exists, and startup links it:")
    configure(False)
    client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "2512-061"})
    check("without a token: not_configured",
          client.get(f"/api/projects/{HIGHLAND}").json()["monday"]["status"] == "not_configured")
    configure(True)
    check("with a token, before any lookup: unchecked, not a stale 'not configured'",
          client.get(f"/api/projects/{HIGHLAND}").json()["monday"]["status"] == "unchecked")
    project_links.refresh_waiting_projects_on_startup()
    p = client.get(f"/api/projects/{HIGHLAND}").json()
    check("the startup refresh links it", p["monday"]["status"] == "linked"
          and p["links"]["monday_board"] == "https://castillope.monday.com/boards/778")

    print("An outage during a save is recorded as an error, not 'not on monday':")
    configure(True)
    monday.portfolio()                                  # warm cache without 999-123
    FAKE["fail"] = "down"
    p = client.patch(f"/api/projects/{HIGHLAND}", json={"castillo_project_id": "999-123"}).json()
    check("unmatched in the cached copy + failed refetch -> error",
          p["castillo_project_id"] == "999-123" and p["monday"]["status"] == "error")
    r = client.post(f"/api/projects/{HIGHLAND}/monday-refresh")
    check("the Retry button reports the outage too", r.json()["monday"]["status"] == "error")
    r = client.post("/api/monday/refresh-projects")
    check("the bulk refresh refuses to stamp an outage onto projects -> 502", r.status_code == 502)
    FAKE["fail"] = None
    configure(False)

    print("Suggestions come from the title block and from monday names:")
    configure(True)
    client.patch(f"/api/projects/{AURORA}", json={"castillo_project_id": None})
    s = client.get(f"/api/projects/{AURORA}/project-id-suggestions").json()["suggestions"]
    check("title block read from the stored PDF", s and s[0]["source"] == "title_block"
          and s[0]["castillo_project_id"] == "2512-053" and s[0]["on_monday"] is True
          and s[0]["name"] == "Aurora 1 IFC")
    check("the same Project ID is not repeated as a name match",
          [x["castillo_project_id"] for x in s].count("2512-053") == 1)
    s = client.get(f"/api/projects/{HIGHLAND}/project-id-suggestions").json()["suggestions"]
    check("stored detection used when the PDF has none",
          s and s[0]["castillo_project_id"] == "254-325")
    check("a similar monday name is offered", any(x["source"] == "name" and
          x["name"] == "Highland North (1 & 2)" for x in s) or s[0]["name"] == "Highland North (1 & 2)")
    listed = {x["id"]: x for x in client.get("/api/projects").json()}
    check("listing exposes what the drawings print",
          listed[HIGHLAND]["title_block_project_id"] == "254-325")

    picker = client.get("/api/monday/portfolio").json()
    check("picker lists items that have a Project ID, by name", picker["configured"] is True
          and [i["castillo_project_id"] for i in picker["items"]] == ["2512-053", "2512-061", "254-325"])

    print("The upload form can carry a Project ID:")
    pdf = io.BytesIO()
    d = rotated_sheet(TITLE_BLOCK)
    d.save(pdf)
    d.close()

    def upload(name, pid):
        return client.post("/api/analyze", data={"project_name": name, "castillo_project_id": pid},
                           files={"file": ("sheet.pdf", pdf.getvalue(), "application/pdf")})

    r = upload("Brand New Project", "2512–061")
    new_pid = r.json()["project_id"]
    p = db.get_project(new_pid)
    check("a new project takes the Project ID", r.status_code == 200 and p["number"] == "2512-061")
    check("monday resolved after the response", p["monday_status"] == "linked"
          and p["monday_board_id"] == "778")
    check("no conflict reported", r.json()["castillo_project_id_conflict"] is None)

    r = upload("Brand New Project", "111-111")
    check("a different Project ID for an existing project is reported, not applied",
          r.json()["castillo_project_id_conflict"]
          == {"castillo_project_id": "2512-061", "submitted": "111-111"}
          and db.get_project(new_pid)["number"] == "2512-061")

    before = {x["id"] for x in client.get("/api/projects").json()}
    uploads_before = len(list(RUNS_DIR.iterdir()))
    r = upload("Should Not Exist", "x" * 60)
    check("an invalid Project ID is rejected before anything is written", r.status_code == 422
          and {x["id"] for x in client.get("/api/projects").json()} == before
          and len(list(RUNS_DIR.iterdir())) == uploads_before)

    print("The export and the chat copilot carry the Project ID:")
    client.patch(f"/api/projects/{AURORA}", json={"castillo_project_id": "2512-053"})
    r = client.get(f"/api/export/{AURORA_RUN}")
    check("export -> 200", r.status_code == 200)
    # Starlette percent-encodes a filename with spaces (filename*=utf-8''...).
    from urllib.parse import unquote
    disposition = unquote(r.headers.get("content-disposition", ""))
    check("filename starts with the Project ID",
          "2512-053_Aurora 1 IFC 60 planset_" in disposition)
    ws = load_workbook(io.BytesIO(r.content))["Summary"]
    check("Summary sheet row 6", ws["A6"].value == "Project ID" and ws["B6"].value == "2512-053")
    from app.exporter import build_workbook
    wb = build_workbook({**db.get_run(AURORA_RUN), "project_name": '=HYPERLINK("https://x","y")'})
    check("a formula-looking project name is exported as text",
          wb["Summary"]["B3"].data_type == "s" and wb["Summary"]["B6"].data_type == "s")

    from app.chat import build_context_pack
    check("chat grounding names the Project ID", "Castillo Project ID: 2512-053"
          in build_context_pack(db.get_run(AURORA_RUN)))

print()
if _FAILS:
    print(f"FAILED ({len(_FAILS)}): {_FAILS}")
    sys.exit(1)
print("ALL PROJECT LINK CHECKS PASSED")
